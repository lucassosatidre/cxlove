#!/usr/bin/env python3
"""Gera o catálogo fechado do salão a partir da planilha mestre auditada."""

import argparse
import json
import unicodedata
from pathlib import Path

from openpyxl import load_workbook


def norm(value):
    text = unicodedata.normalize("NFD", str(value or ""))
    text = "".join(c for c in text if unicodedata.category(c) != "Mn")
    return " ".join(text.lower().split())


def tipo_cozinha(tipo):
    tipo = norm(tipo)
    if tipo.startswith("sabor doce"):
        return "sabor_doce"
    if tipo.startswith("sabor"):
        return "sabor"
    if tipo == "adicional":
        return "adicional"
    if tipo == "borda":
        return "borda"
    if tipo == "dip":
        return "dip"
    if tipo in ("bebida", "open bar"):
        return "bebida"
    return "outro"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("mestre", type=Path)
    ap.add_argument("saida", type=Path)
    args = ap.parse_args()

    ws = load_workbook(args.mestre, read_only=True, data_only=True)["Mapa Integração"]
    rows = ws.iter_rows(min_row=4, values_only=True)
    headers = list(next(rows))
    ix = {name: pos for pos, name in enumerate(headers)}
    por_codigo = {}
    por_texto = {}
    ambiguos = {}

    for row in rows:
        if row[ix["Canal"]] != "Salão":
            continue
        nome_saipos = str(row[ix["Nome Saipos"]] or "").strip()
        nome_padrao = str(row[ix["Nome Padrão"]] or "").strip()
        codigo = str(row[ix["Código Opção"]] or "").strip()
        item = {"tipo": tipo_cozinha(row[ix["Tipo Opção"]]), "nome": nome_padrao}
        if codigo:
            por_codigo[codigo] = item
            por_codigo[codigo.split(".")[-1]] = item
        if nome_saipos:
            chave = norm(nome_saipos)
            anterior = por_texto.get(chave)
            if anterior and anterior != item:
                ambiguos.setdefault(chave, [anterior])
                if item not in ambiguos[chave]:
                    ambiguos[chave].append(item)
            else:
                por_texto[chave] = item
        if nome_padrao:
            chave_padrao = norm(nome_padrao)
            if chave_padrao not in por_texto:
                por_texto[chave_padrao] = item

    for chave in ambiguos:
        por_texto.pop(chave, None)

    payload = {
        "versao": 1,
        "fonte": "cardapio_mestre_mana_provisao_saipos_v2.xlsx",
        "por_codigo": dict(sorted(por_codigo.items())),
        "por_texto": dict(sorted(por_texto.items())),
        "textos_ambiguos": dict(sorted(ambiguos.items())),
    }
    args.saida.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"codigos": len(por_codigo), "textos": len(por_texto), "ambiguos": len(ambiguos)}))


if __name__ == "__main__":
    main()
